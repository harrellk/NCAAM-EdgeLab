import pandas as pd
import numpy as np
from pathlib import Path
import sys

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


EP_BINS = [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, np.inf]
EP_LABELS = [
    "0–0.99",
    "1–1.99",
    "2–2.99",
    "3–3.99",
    "4–4.99",
    "5–5.99",
    "6–6.99",
    "7–7.99",
    "8–8.99",
    "9+",
]


# ---------------------------------------------------------
# CLASSIFICATION RULES
# ---------------------------------------------------------
def classify_group(row):
    is_neutral = str(row["IsNeutral"]).lower() == "true"
    edge_side = str(row["EdgeSide"]).upper()

    if is_neutral:
        return "Neutral"
    if not is_neutral and edge_side == "AWAY":
        return "NonNeutral_AwayEdge"
    if not is_neutral and edge_side == "HOME":
        return "NonNeutral_HomeEdge"
    return "Other"


# ---------------------------------------------------------
# BUILD EDGE BUCKET TABLE
# ---------------------------------------------------------
def build_edge_bucket_table(df):
    d = df.copy()
    d["EP_Bucket"] = pd.cut(
        d["EdgePoints"], bins=EP_BINS, labels=EP_LABELS, right=False
    )
    d["GroupCat"] = d.apply(classify_group, axis=1)

    g = d.groupby(["EP_Bucket", "GroupCat"])["EdgeHit"].agg(["count", "sum"])
    g["WinRate"] = g["sum"] / g["count"]

    # Format WinRate as xx.xx%
    g["WinRate"] = (g["WinRate"] * 100).round(2).astype(str) + "%"

    return g.reset_index()


# ---------------------------------------------------------
# EXCEL STYLING FUNCTION
# ---------------------------------------------------------
def style_sheet(ws):
    """Apply soft/clean formatting to an openpyxl worksheet."""

    # Color palette
    header_fill = PatternFill(
        start_color="DCE6F1", end_color="DCE6F1", fill_type="solid"
    )
    alt_fill = PatternFill(start_color="F7F7F7", end_color="F7F7F7", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    # Header formatting
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Row striping + borders
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in row:
            if i % 2 == 0:
                cell.fill = alt_fill
            cell.border = thin_border

            # Center WinRate column
            if cell.column_letter == "E":  # WinRate is col 5
                cell.alignment = Alignment(horizontal="center")

    # Auto-fit columns
    for col in ws.columns:
        max_len = max(
            len(str(cell.value)) if cell.value is not None else 0 for cell in col
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 3


# ---------------------------------------------------------
# GENERATE REPORTS
# ---------------------------------------------------------
def generate_edge_bucket_reports(results_csv, report_date, out_excel):
    df = pd.read_csv(results_csv)
    df["Date_dt"] = pd.to_datetime(df["Date"])
    target_dt = pd.to_datetime(report_date)

    # Daily slice
    df_daily = df[df["Date_dt"] == target_dt]

    # Rollup slice (start = 2025-11-22)
    start_dt = pd.to_datetime("2025-11-22")
    df_rollup = df[(df["Date_dt"] >= start_dt) & (df["Date_dt"] <= target_dt)]

    # Entire history
    df_full = df

    # Build tables
    t_daily = build_edge_bucket_table(df_daily)
    t_rollup = build_edge_bucket_table(df_rollup)
    t_full = build_edge_bucket_table(df_full)

    # Ensure output directory exists
    out_excel = Path(out_excel)
    out_excel.parent.mkdir(parents=True, exist_ok=True)

    # Write Excel with styling
    with pd.ExcelWriter(out_excel, engine="openpyxl") as writer:
        t_daily.to_excel(writer, sheet_name=f"Daily_{report_date}", index=False)
        t_rollup.to_excel(writer, sheet_name=f"Rollup_to_{report_date}", index=False)
        t_full.to_excel(writer, sheet_name="Master_All", index=False)

        # Apply styling
        wb = writer.book
        ws_daily = wb[f"Daily_{report_date}"]
        ws_roll = wb[f"Rollup_to_{report_date}"]
        ws_full = wb["Master_All"]

        style_sheet(ws_daily)
        style_sheet(ws_roll)
        style_sheet(ws_full)

    return out_excel


# ---------------------------------------------------------
# COMMAND-LINE ENTRY POINT
# ---------------------------------------------------------
if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("\nUsage:")
        print(
            "python report_edge_buckets.py <results_csv> <YYYY-MM-DD> <output_excel_file>"
        )
        print("\nExample:")
        print(
            "python report_edge_buckets.py output/reports/model_results_tracking.csv 2025-11-26 output/reports/edge_bucket_report_2025-11-26.xlsx"
        )
        sys.exit(1)

    results_csv = sys.argv[1]
    report_date = sys.argv[2]
    out_excel = sys.argv[3]

    path = generate_edge_bucket_reports(results_csv, report_date, out_excel)
    print(f"\n✔ Pretty Excel report created: {path}")
