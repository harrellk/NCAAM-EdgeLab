import pandas as pd
import numpy as np
from pathlib import Path
from openpyxl import Workbook

# =====================================================
# CONFIG LOCATIONS
# =====================================================
RESULTS_FILE = Path("output/reports/model_results_tracking.csv")
OUTPUT_FILE = Path("output/reports/team_performance_report_v2.xlsx")


# =====================================================
# LOAD DATA
# =====================================================
def load_results():
    if not RESULTS_FILE.exists():
        raise FileNotFoundError(f"Missing results tracking file: {RESULTS_FILE}")
    return pd.read_csv(RESULTS_FILE)


# =====================================================
# TEAM-WIDE ATS PERFORMANCE (True ATS)
# =====================================================
def build_team_wide_ats(df):
    df["HomeATS"] = (df["ActualMargin"] + df["HomeMarketSpread"]) > 0
    df["AwayATS"] = ((-df["ActualMargin"]) + df["AwayMarketSpread"]) > 0

    home = (
        df.groupby("HomeTeam")
        .agg(Games=("HomeATS", "count"), ATS_Wins=("HomeATS", "sum"))
        .reset_index()
        .rename(columns={"HomeTeam": "Team"})
    )

    away = (
        df.groupby("AwayTeam")
        .agg(Games=("AwayATS", "count"), ATS_Wins=("AwayATS", "sum"))
        .reset_index()
        .rename(columns={"AwayTeam": "Team"})
    )

    merged = home.merge(
        away, on="Team", how="outer", suffixes=("_home", "_away")
    ).fillna(0)
    merged["Games"] = merged["Games_home"] + merged["Games_away"]
    merged["ATS_Wins"] = merged["ATS_Wins_home"] + merged["ATS_Wins_away"]
    merged["ATS_HitRate"] = merged["ATS_Wins"] / merged["Games"]
    return merged[["Team", "Games", "ATS_Wins", "ATS_HitRate"]].sort_values("Team")


# =====================================================
# MODEL ACCURACY (Using ALL games, normalized)
# =====================================================
def build_model_accuracy(df):
    rows = []

    for _, row in df.iterrows():
        # TEAM_A perspective
        A = row["Team_A"]
        model_margin_A = row["ModelMargin"]
        actual_margin_A = row["ActualMargin"]
        error_A = model_margin_A - actual_margin_A

        rows.append({"Team": A, "SpreadError": error_A})

        # TEAM_B perspective = flip signs
        B = row["Team_B"]
        model_margin_B = -row["ModelMargin"]
        actual_margin_B = -row["ActualMargin"]
        error_B = model_margin_B - actual_margin_B

        rows.append({"Team": B, "SpreadError": error_B})

    # Build dataframe of normalized errors
    df_norm = pd.DataFrame(rows)
    df_norm["AbsError"] = df_norm["SpreadError"].abs()

    group = (
        df_norm.groupby("Team")
        .agg(
            Games=("SpreadError", "count"),
            MAE=("AbsError", "mean"),
            RMSE=("SpreadError", lambda x: np.sqrt((x**2).mean())),
            Bias=("SpreadError", "mean"),
        )
        .reset_index()
    )

    return group


# =====================================================
# EDGE PERFORMANCE (When model picks that team)
# =====================================================
def build_edge_performance(df):
    group = (
        df.groupby("EdgeTeam")
        .agg(
            Games=("EdgeHit", "count"),
            ATS_Wins=("EdgeHit", "sum"),
            ATS_HitRate=("EdgeHit", "mean"),
            AvgEdgePoints=("EdgePoints", "mean"),
        )
        .reset_index()
    )
    return group


# =====================================================
# SAVE ALL REPORTS
# =====================================================
def write_excel(ats, acc, edge):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Team_Wide_ATS"
    ws1.append(list(ats.columns))
    for row in ats.itertuples(index=False):
        ws1.append(list(row))

    ws2 = wb.create_sheet("Model_Accuracy")
    ws2.append(list(acc.columns))
    for row in acc.itertuples(index=False):
        ws2.append(list(row))

    ws3 = wb.create_sheet("Edge_Performance")
    ws3.append(list(edge.columns))
    for row in edge.itertuples(index=False):
        ws3.append(list(row))

    wb.save(OUTPUT_FILE)
    print(f"🔥 Team Performance Report updated: {OUTPUT_FILE}")


# =====================================================
# MAIN
# =====================================================
def main():
    df = load_results()

    ats = build_team_wide_ats(df)
    acc = build_model_accuracy(df)
    edge = build_edge_performance(df)

    write_excel(ats, acc, edge)


if __name__ == "__main__":
    main()
